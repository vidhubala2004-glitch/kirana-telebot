import os
from pathlib import Path
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from pptx import Presentation
from pptx.util import Inches
from config import BASE_DIR, INVOICE_DIR, REPORT_DIR
from database import get_db

os.environ.setdefault("MPLCONFIGDIR", str(BASE_DIR / "data" / ".matplotlib"))
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


def get_bill_for_invoice(bill_id):
    with get_db() as db:
        bill=db.execute("SELECT * FROM bills WHERE id=? AND status='FINALIZED'",(int(bill_id),)).fetchone()
        if not bill: return None
        items=db.execute("SELECT bi.*,p.name,p.unit,p.hsn_code FROM bill_items bi JOIN products p ON p.id=bi.product_id WHERE bi.bill_id=? ORDER BY bi.id",(int(bill_id),)).fetchall()
    return bill,items


def generate_invoice_pdf(bill_id):
    data=get_bill_for_invoice(bill_id)
    if not data: return {"success":False,"error":"Finalized bill not found."}
    bill,items=data
    file_path=INVOICE_DIR/f"{bill['bill_number']}.pdf"
    doc=SimpleDocTemplate(str(file_path),pagesize=A4,rightMargin=12*mm,leftMargin=12*mm,topMargin=12*mm,bottomMargin=12*mm)
    styles=getSampleStyleSheet(); story=[]
    story += [Paragraph("SRI LAKSHMI KIRANA STORES",styles["Title"]),Paragraph("GST INVOICE",styles["Heading2"]),Spacer(1,8)]
    story += [Paragraph(f"<b>Invoice No:</b> {bill['bill_number']}",styles["Normal"]),Paragraph(f"<b>Date:</b> {bill['created_at']}",styles["Normal"]),Paragraph(f"<b>Customer:</b> {bill['customer_name'] or 'Walk-in Customer'}",styles["Normal"]),Paragraph(f"<b>Payment:</b> {bill['payment_mode']}",styles["Normal"])]
    if bill["payment_reference"]: story.append(Paragraph(f"<b>Reference:</b> {bill['payment_reference']}",styles["Normal"]))
    story.append(Spacer(1,10))
    data_rows=[["Product","HSN","Qty","Unit Price","GST","CGST","SGST","Total"]]
    for x in items:
        data_rows.append([x["name"],x["hsn_code"] or "-",str(x["quantity"]),f"₹{x['unit_price']/100:.2f}",f"{x['gst_rate']}%",f"₹{x['cgst']/100:.2f}",f"₹{x['sgst']/100:.2f}",f"₹{x['total']/100:.2f}"])
    table=Table(data_rows,repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("GRID",(0,0),(-1,-1),0.5,colors.grey),("ALIGN",(2,1),(-1,-1),"RIGHT"),("VALIGN",(0,0),(-1,-1),"MIDDLE")]))
    story += [table,Spacer(1,12)]
    totals=[["Subtotal",f"₹{bill['subtotal']/100:.2f}"],["Discount",f"₹{(bill['discount'] or 0)/100:.2f}"],["CGST",f"₹{bill['cgst']/100:.2f}"],["SGST",f"₹{bill['sgst']/100:.2f}"],["Total Tax",f"₹{bill['total_tax']/100:.2f}"],["FINAL AMOUNT",f"₹{bill['grand_total']/100:.2f}"]]
    tt=Table(totals,colWidths=[100*mm,50*mm],hAlign="RIGHT")
    tt.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.5,colors.grey),("ALIGN",(1,0),(1,-1),"RIGHT"),("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold")]))
    story += [tt,Spacer(1,20),Paragraph("Thank you for shopping with us!",styles["Normal"])]
    doc.build(story)
    return {"success":True,"bill_id":bill_id,"bill_number":bill["bill_number"],"file_path":str(file_path)}


def generate_sales_deck():
    from tools.analytics import weekly_sales
    report=weekly_sales(); file_path=REPORT_DIR/"weekly_sales_analysis.pptx"; chart_path=REPORT_DIR/"weekly_sales_chart.png"
    dates=[d["date"] for d in report["days"]]; sales=[d["total"] for d in report["days"]]
    plt.figure(figsize=(10,5)); plt.plot(dates,sales,marker="o"); plt.title("Weekly Sales"); plt.xlabel("Date"); plt.ylabel("Sales (₹)"); plt.xticks(rotation=45); plt.tight_layout(); plt.savefig(chart_path); plt.close()
    prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text="Kirana Store Weekly Analysis"; slide.placeholders[1].text=f"{report['start_date']} to {report['end_date']}"
    slide=prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text="Weekly Summary"; avg=report["total_sales"]/report["total_bills"] if report["total_bills"] else 0; slide.placeholders[1].text=f"Total Sales: ₹{report['total_sales']:.2f}\n\nTotal Bills: {report['total_bills']}\n\nAverage Bill: ₹{avg:.2f}"
    slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text="Daily Sales Trend"; slide.shapes.add_picture(str(chart_path),Inches(.7),Inches(1.4),width=Inches(8.5))
    slide=prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text="Top Products"; totals={}
    for d in report["days"]:
        for p in d["top_products"]: totals[p["name"]]=totals.get(p["name"],0)+p["revenue"]
    ranked=sorted(totals.items(),key=lambda x:x[1],reverse=True); slide.placeholders[1].text="\n".join(f"{n}: ₹{r:.2f}" for n,r in ranked[:5]) or "No sales data."
    slide=prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text="Business Insights"; slide.placeholders[1].text=f"Weekly revenue: ₹{report['total_sales']:.2f}\n\nBills: {report['total_bills']}\n\nAverage bill: ₹{avg:.2f}\n\nUse top-selling products to plan future stock purchases."
    prs.save(file_path)
    return {"success":True,"file_path":str(file_path),"message":"Weekly sales PowerPoint generated."}
